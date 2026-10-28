import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# PRESENTATION & PALETTE CONSTANTS (Faithfully matching 1month ppt.pptx)
# -----------------------------------------------------------------------------
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

COLOR_CRIMSON = RGBColor(225, 29, 72)       # Primary Crimson Red (#E11D48)
COLOR_CRIMSON_DARK = RGBColor(190, 18, 60)  # Deep Crimson (#BE123C)
COLOR_CRIMSON_LIGHT = RGBColor(254, 226, 226)# Border tint (#FEE2E2)
COLOR_PINK_TINT = RGBColor(255, 241, 242)   # Soft Rose Surface (#FFF1F2)
COLOR_WHITE = RGBColor(255, 255, 255)       # Card Surface
COLOR_TEXT_DARK = RGBColor(15, 23, 42)      # Charcoal Slate (#0F172A)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)  # Slate Grey (#64748B)
COLOR_TEXT_DIM = RGBColor(148, 163, 184)    # Light Slate (#94A3B8)
COLOR_CARD_BORDER = RGBColor(254, 205, 211) # Rose 200 border (#FECDD3)
COLOR_BORDER_SUBTLE = RGBColor(226, 232, 240)# Neutral subtle border (#E2E8F0)

COLOR_GREEN_BG = RGBColor(209, 250, 229)    # Emerald 100
COLOR_GREEN_TEXT = RGBColor(5, 150, 105)    # Emerald 600
COLOR_BLUE_BG = RGBColor(219, 234, 254)     # Blue 100
COLOR_BLUE_TEXT = RGBColor(37, 99, 235)     # Blue 600
COLOR_AMBER_BG = RGBColor(254, 243, 199)    # Amber 100
COLOR_AMBER_TEXT = RGBColor(217, 119, 6)    # Amber 600

FONT_DISPLAY = 'Georgia'
FONT_BODY = 'Segoe UI'
FONT_MONO = 'Consolas'

NIET_LOGO_PATH = os.path.abspath('extracted_1month_ppt/niet_logo_trans.png')

# -----------------------------------------------------------------------------
# CORE SLIDE GENERATOR CLASS
# -----------------------------------------------------------------------------
class JavaPBLPresentationBuilder:
    def __init__(self):
        self.prs = pptx.Presentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT
        self.blank_layout = self.prs.slide_layouts[6] # completely blank layout

    def add_slide(self):
        return self.prs.slides.add_slide(self.blank_layout)

    def draw_header(self, slide, num_str, title_w1, title_w2, subtitle="ACADEMIC PROJECT TRACKING SYSTEM"):
        """Renders standard top header with pill badge, serif heading, slogan, and NIET logo."""
        # 1. Number Pill Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.35), Inches(0.85), Inches(0.85))
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_PINK_TINT
        badge.line.color.rgb = COLOR_CARD_BORDER
        badge.line.width = Pt(1.5)
        tf = badge.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = num_str
        p.font.name = FONT_DISPLAY
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_CRIMSON
        p.alignment = PP_ALIGN.CENTER

        # 2. Main Title Text Box
        t_box = slide.shapes.add_textbox(Inches(1.6), Inches(0.32), Inches(7.6), Inches(0.9))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_top = Inches(0)
        tf_t.margin_left = Inches(0)
        p_title = tf_t.paragraphs[0]
        
        # Word 1 (Dark)
        r1 = p_title.add_run()
        r1.text = title_w1 + " "
        r1.font.name = FONT_DISPLAY
        r1.font.size = Pt(26)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_DARK
        
        # Word 2 (Crimson)
        r2 = p_title.add_run()
        r2.text = title_w2
        r2.font.name = FONT_DISPLAY
        r2.font.size = Pt(26)
        r2.font.bold = True
        r2.font.color.rgb = COLOR_CRIMSON

        # Subtitle
        p_sub = tf_t.add_paragraph()
        r_sub = p_sub.add_run()
        r_sub.text = subtitle
        r_sub.font.name = FONT_BODY
        r_sub.font.size = Pt(9.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = COLOR_TEXT_MUTED

        # Thin accent line under title
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.6), Inches(1.22), Inches(4.5), Pt(1.5))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_CRIMSON
        line.line.fill.background()

        # 3. Top Right Institutional Banner
        top_r = slide.shapes.add_textbox(Inches(8.5), Inches(0.4), Inches(2.8), Inches(0.4))
        tf_r = top_r.text_frame
        tf_r.margin_top = Inches(0)
        tf_r.margin_right = Inches(0)
        p_r = tf_r.paragraphs[0]
        p_r.alignment = PP_ALIGN.RIGHT
        r_r = p_r.add_run()
        r_r.text = "— LEARN  |  INNOVATE  |  TRANSFORM"
        r_r.font.name = FONT_BODY
        r_r.font.size = Pt(9.5)
        r_r.font.bold = True
        r_r.font.color.rgb = COLOR_TEXT_MUTED

        # 4. NIET Logo
        if os.path.exists(NIET_LOGO_PATH):
            slide.shapes.add_picture(NIET_LOGO_PATH, Inches(11.5), Inches(0.28), Inches(1.3), Inches(0.55))

        # 5. Footer Bar
        # Horizontal divider
        f_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.9), Inches(12.13), Pt(1))
        f_line.fill.solid()
        f_line.fill.fore_color.rgb = COLOR_BORDER_SUBTLE
        f_line.line.fill.background()

        # Left footer label
        f_box_l = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(5.0), Inches(0.35))
        p_fl = f_box_l.text_frame.paragraphs[0]
        r_fl = p_fl.add_run()
        r_fl.text = "Department of CSE (AI)  |  Session 2026–27  |  Capstone Project"
        r_fl.font.name = FONT_BODY
        r_fl.font.size = Pt(9)
        r_fl.font.color.rgb = COLOR_TEXT_MUTED

        # Right footer page number
        f_box_r = slide.shapes.add_textbox(Inches(11.0), Inches(6.92), Inches(1.73), Inches(0.4))
        p_fr = f_box_r.text_frame.paragraphs[0]
        p_fr.alignment = PP_ALIGN.RIGHT
        r_fr = p_fr.add_run()
        r_fr.text = f"{num_str}  —"
        r_fr.font.name = FONT_BODY
        r_fr.font.size = Pt(12)
        r_fr.font.bold = True
        r_fr.font.color.rgb = COLOR_CRIMSON

    def add_card(self, slide, left, top, width, height, bg_rgb=COLOR_WHITE, border_rgb=COLOR_CARD_BORDER, border_pt=1):
        """Creates a standardized rounded container card."""
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_rgb
        if border_rgb:
            shape.line.color.rgb = border_rgb
            shape.line.width = Pt(border_pt)
        else:
            shape.line.fill.background()
        return shape

    def add_bottom_bar(self, slide, title, text, icon="🏆"):
        """Creates signature bottom summary / takeaway card."""
        card = self.add_card(slide, Inches(0.6), Inches(6.05), Inches(12.13), Inches(0.72), bg_rgb=COLOR_PINK_TINT, border_rgb=COLOR_CARD_BORDER)
        
        # Icon badge
        ibadge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(6.15), Inches(0.52), Inches(0.52))
        ibadge.fill.solid()
        ibadge.fill.fore_color.rgb = COLOR_WHITE
        ibadge.line.color.rgb = COLOR_CARD_BORDER
        ibadge.line.width = Pt(1)
        tf_i = ibadge.text_frame
        tf_i.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_i = tf_i.paragraphs[0]
        p_i.text = icon
        p_i.font.size = Pt(14)
        p_i.alignment = PP_ALIGN.CENTER

        # Text box
        tbox = slide.shapes.add_textbox(Inches(1.4), Inches(6.1), Inches(11.1), Inches(0.6))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        
        r_t = p.add_run()
        r_t.text = title + "  "
        r_t.font.name = FONT_BODY
        r_t.font.size = Pt(11)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_CRIMSON
        
        r_d = p.add_run()
        r_d.text = text
        r_d.font.name = FONT_BODY
        r_d.font.size = Pt(10.5)
        r_d.font.color.rgb = COLOR_TEXT_DARK

    def add_pill_chip(self, slide, left, top, text, bg_rgb, text_rgb, width=None, height=Inches(0.32)):
        """Creates small colored badge chip."""
        w = width if width else Inches(len(text) * 0.09 + 0.3)
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, height)
        chip.fill.solid()
        chip.fill.fore_color.rgb = bg_rgb
        chip.line.fill.background()
        tf = chip.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = FONT_BODY
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = text_rgb
        p.alignment = PP_ALIGN.CENTER
        return chip

    def save(self, filepath):
        self.prs.save(filepath)
        print(f"[+] Presentation saved successfully: {filepath}")

print("JavaPBLPresentationBuilder base class ready.")
